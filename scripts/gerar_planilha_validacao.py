"""Gera validacao/casos_validacao.xlsx: casos reais para o contador conferir o motor.

Uso: python scripts/gerar_planilha_validacao.py [caminho_saida.xlsx]

Cada caso lista os valores calculados pelo motor; o contador preenche a coluna
"Valor correto" e a planilha mostra a diferença. A aba "Memória de cálculo"
traz cada passo com fórmula e base legal.
"""

from __future__ import annotations

import sys
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from tolaris_motor import Atividade, PerfilEmpresa, Resultado, brl, pct
from tolaris_motor import pessoa_fisica as pf
from tolaris_motor import autonomo, avisos, enquadramento, presumido, real, recuperacao, reforma, reforma_simples, simples

FONTE = "Arial"
AZUL_ESCURO = "1F3A5F"
AMARELO = PatternFill("solid", start_color="FFF2CC")
CINZA = PatternFill("solid", start_color="F2F2F2")
CABECALHO = PatternFill("solid", start_color=AZUL_ESCURO)
BORDA = Border(bottom=Side(style="thin", color="D9D9D9"))
MOEDA = '"R$" #,##0.00;-"R$" #,##0.00;"-"'
JAN_2026 = date(2026, 1, 1)


@dataclass
class Caso:
    codigo: str
    modulo: str
    descricao: str
    resultado: Resultado


def _perfil_desc(p: PerfilEmpresa) -> str:
    partes = [f"{p.atividade.value}", f"receita {brl(p.receita_bruta_anual)}/ano"]
    for campo, rotulo in (
        ("folha_clt_anual", "folha CLT"),
        ("pro_labore_anual", "pró-labore"),
        ("compras_anual", "compras"),
        ("despesas_creditaveis_anual", "desp. creditáveis"),
        ("outras_despesas_anual", "outras desp."),
    ):
        v = getattr(p, campo)
        if v:
            partes.append(f"{rotulo} {brl(v)}")
    partes.append(f"ano {p.ano}")
    if p.atividade.vende_mercadoria:
        partes.append(f"ICMS {pct(p.aliquota_icms, 0)}")
    else:
        partes.append(f"ISS {pct(p.aliquota_iss, 0)}")
    return "; ".join(partes)


def _resultado_enquadramento(p: PerfilEmpresa, percentual: str | None = None) -> Resultado:
    """Um Resultado por empresa: o DAS de cada caminho (hoje, pró-labore até o fator R de 28% e parte no Anexo III).

    As linhas de ``tributos`` são alternativas, não se somam; a memória traz o fator R e a alíquota efetiva de cada caminho.
    """
    r = Resultado("Enquadramento")
    for c in enquadramento.analisar(p, percentual):
        simples_res = next(x for x in c.comparativo.resultados if x.titulo.startswith("Simples"))
        r.tributos[f"DAS - {c.descricao}"] = simples_res.tributos["DAS"] if simples_res.aplicavel else Decimal(0)
        if c.custo_extra_anual:
            r.tributos["INSS do sócio sobre o pró-labore adicional (caminho do fator R)"] = c.custo_extra_anual
        for passo in simples_res.memoria:
            r.passo(f"{c.chave}: {passo.descricao}", passo.formula, passo.valor, passo.base_legal)
    r.alerta("Os caminhos são alternativas: o DAS de cada um é calculado como se a empresa seguisse só aquele caminho.")
    r.alerta("Receita no Anexo III só vale se a empresa exerce de fato a atividade, tem o CNAE no cadastro e emite a nota com ela "
             "(LC 123/2006, art. 18, §§ 4º e 5º-B). O motor só mostra o efeito no imposto; a decisão é do contador.")
    r.alerta("O INSS do sócio (11% sobre o pró-labore a mais, dentro do teto) é custo do caminho do fator R; o IRRF não entra.")
    return r


def casos() -> list[Caso]:
    lista: list[Caso] = []

    perfis = {
        "E1": PerfilEmpresa(Atividade.SERVICO_FATOR_R, 600_000, folha_clt_anual=96_000, pro_labore_anual=60_000),
        "E2": PerfilEmpresa(Atividade.SERVICO_FATOR_R, 600_000, folha_clt_anual=96_000, pro_labore_anual=72_000),
        "E3": PerfilEmpresa(Atividade.COMERCIO, 2_000_000, folha_clt_anual=240_000, pro_labore_anual=36_000,
                            compras_anual=1_200_000, despesas_creditaveis_anual=80_000, outras_despesas_anual=150_000),
        "E4": PerfilEmpresa(Atividade.INDUSTRIA, 3_000_000, folha_clt_anual=600_000, pro_labore_anual=60_000,
                            compras_anual=1_500_000, despesas_creditaveis_anual=200_000, outras_despesas_anual=100_000),
        "E5": PerfilEmpresa(Atividade.SERVICO_ANEXO_IV, 900_000, folha_clt_anual=300_000, pro_labore_anual=36_000,
                            outras_despesas_anual=200_000, aliquota_iss="0.03"),
        "E6": PerfilEmpresa(Atividade.SERVICO_ANEXO_III, 8_000_000, folha_clt_anual=1_500_000,
                            pro_labore_anual=120_000, outras_despesas_anual=2_000_000),
        "E7": PerfilEmpresa(Atividade.SERVICO_FATOR_R, 1_200_000, folha_clt_anual=150_000, pro_labore_anual=48_000,
                            outras_despesas_anual=300_000, despesas_creditaveis_anual=60_000),
    }
    for cod, p in perfis.items():
        for mod in (simples, presumido, real):
            r = mod.calcular(p)
            sigla = {"Simples Nacional": "SN", "Lucro Presumido": "LP", "Lucro Real": "LR"}[r.titulo]
            lista.append(Caso(f"{cod}-{sigla}", r.titulo, _perfil_desc(p), r))

    for i, (rend, dep) in enumerate(((3_500, 0), (5_000, 0), (5_900, 0), (7_000, 1), (12_000, 2)), start=1):
        inss = pf.inss_empregado(rend, JAN_2026)
        r = pf.irpf_mensal(rend, JAN_2026, inss=inss, dependentes=dep)
        r.tributos = {"INSS empregado": inss, **r.tributos}
        lista.append(Caso(f"PF{i}", "IRPF/INSS mensal (jan/2026)",
                          f"salário {brl(rend)}; {dep} dependente(s); CLT", r))

    for i, (retirada, regime) in enumerate(((15_000, "presumido"), (30_000, "real"), (80_000, "presumido")), start=1):
        for cen in pf.comparar_pro_labore_dividendos(retirada, JAN_2026, regime=regime):
            lista.append(Caso(f"SOC{i}-{'A' if cen.dividendos == 0 else 'B'}", f"Pró-labore x dividendos ({regime})",
                              f"retirada {brl(retirada)}/mês; {cen.descricao}", cen.resultado))

    entradas = {
        "RT1": ("Comércio no Presumido",
                reforma.EntradaReforma(2_000_000, vende_mercadoria=True, compras_anual=1_200_000)),
        "RT2": ("Serviços (consultoria) no Presumido",
                reforma.EntradaReforma(1_200_000, vende_mercadoria=False, despesas_creditaveis_anual=120_000)),
        "RT3": ("Escritório contábil (redução de 30%)",
                reforma.EntradaReforma(600_000, vende_mercadoria=False, despesas_creditaveis_anual=60_000,
                                       reducao_setorial="0.3")),
        "RT4": ("Indústria no Lucro Real",
                reforma.EntradaReforma(3_000_000, vende_mercadoria=True, pis_cofins_nao_cumulativo=True,
                                       compras_anual=1_500_000, despesas_creditaveis_anual=200_000,
                                       aliquota_ipi="0.05")),
    }
    for cod, (desc, e) in entradas.items():
        for a in reforma.simular(e):
            if a.ano in (2025, 2027, 2029, 2033):
                lista.append(Caso(f"{cod}-{a.ano}", f"Reforma - {a.ano}",
                                  f"{desc}; receita {brl(e.receita_bruta_anual)}", a.resultado))

    simples_b2b = {
        "SR1": ("Comércio, 80% B2B, compras de fornecedores do regime regular",
                reforma_simples.EntradaSimplesReforma("comercio", 2_000_000, "0.8", compras_creditaveis_anual=1_200_000)),
        "SR2": ("Serviços (Anexo III), 100% B2B, sem compras",
                reforma_simples.EntradaSimplesReforma("servico_anexo_iii", 600_000, 1)),
        "SR3": ("Serviços com fator R, 30% B2B",
                reforma_simples.EntradaSimplesReforma("servico_fator_r", 1_200_000, "0.3", folha_fator_r_anual=400_000)),
    }
    for cod, (desc, e) in simples_b2b.items():
        for c in reforma_simples.simular(e):
            lista.append(Caso(f"{cod}-{c.ano}", f"Simples na Reforma - {c.ano}",
                              f"{desc}; receita {brl(e.receita_bruta_anual)}", c.resultado))

    enq = {
        "ENQ1": ("Serviços do fator R, pró-labore de R$ 3.000/mês e sem empregados (o caso do kit de teste)",
                 PerfilEmpresa(Atividade.SERVICO_FATOR_R, 600_000, pro_labore_anual=36_000, outras_despesas_anual=60_000), "0.3"),
        "ENQ2": ("Serviços do fator R com um empregado e pró-labore baixo",
                 PerfilEmpresa(Atividade.SERVICO_FATOR_R, 1_200_000, folha_clt_anual=150_000, pro_labore_anual=48_000,
                               outras_despesas_anual=300_000), None),
        "ENQ3": ("Serviços do fator R, receita baixa e só pró-labore",
                 PerfilEmpresa(Atividade.SERVICO_FATOR_R, 300_000, pro_labore_anual=24_000, outras_despesas_anual=30_000), None),
        "ENQ4": ("Serviços do fator R, receita alta (5ª faixa) e fator R de 18,5%",
                 PerfilEmpresa(Atividade.SERVICO_FATOR_R, 2_400_000, folha_clt_anual=300_000, pro_labore_anual=120_000,
                               outras_despesas_anual=600_000), None),
        "ENQ5": ("Fator R já acima de 28% (folha CLT de R$ 140.000, pró-labore de R$ 28.000 e FGTS): só o cenário de hoje",
                 PerfilEmpresa(Atividade.SERVICO_FATOR_R, 600_000, folha_clt_anual=140_000, pro_labore_anual=28_000,
                               outras_despesas_anual=60_000), None),
        "ENQ6": ("Folha do fator R informada pelo contador (R$ 100.000) em vez da calculada, receita de R$ 600.000",
                 PerfilEmpresa(Atividade.SERVICO_FATOR_R, 600_000, pro_labore_anual=36_000, outras_despesas_anual=60_000,
                               folha_fator_r_anual=100_000), None),
    }
    for cod, (desc, p, percentual) in enq.items():
        extra = f"; cenário próprio de {int(Decimal(percentual) * 100)}% no Anexo III" if percentual else ""
        lista.append(Caso(cod, "Enquadramento: Anexo V ao III (2026)", f"{desc}; {_perfil_desc(p)}{extra}",
                          _resultado_enquadramento(p, percentual)))

    simples_b2b_intermediarios = {
        "SR4": ("Comércio, 80% B2B, compras de fornecedores do regime regular",
                reforma_simples.EntradaSimplesReforma("comercio", 2_000_000, "0.8", compras_creditaveis_anual=1_200_000,
                                                      anos=[2028, 2029, 2030, 2031, 2032])),
        "SR5": ("Serviços (Anexo III), 100% B2B, sem compras",
                reforma_simples.EntradaSimplesReforma("servico_anexo_iii", 600_000, 1, anos=[2029, 2031])),
        "SR6": ("Serviços (Anexo III), 5ª faixa (o teto do ISS no DAS atua), 50% B2B",
                reforma_simples.EntradaSimplesReforma("servico_anexo_iii", 3_000_000, "0.5", anos=[2028, 2029, 2032])),
    }
    for cod, (desc, e) in simples_b2b_intermediarios.items():
        for c in reforma_simples.simular(e):
            lista.append(Caso(f"{cod}-{c.ano}", f"Simples na Reforma - {c.ano}",
                              f"{desc}; receita {brl(e.receita_bruta_anual)}", c.resultado))

    profissionais = {
        "AU1": ("Receita R$ 6.000/mês, pode ser MEI", autonomo.EntradaAutonomo(6000, mei_permitido=True, contabilidade_mensal=300)),
        "AU2": ("Receita R$ 15.000/mês, despesas R$ 1.000", autonomo.EntradaAutonomo(15000, despesas_mensais=1000, contabilidade_mensal=400)),
        "AU3": ("Receita R$ 40.000/mês, 1 dependente", autonomo.EntradaAutonomo(40000, dependentes=1, contabilidade_mensal=600)),
    }
    for cod, (desc, e) in profissionais.items():
        for o in autonomo.comparar(e).opcoes:
            if o.aplicavel:
                sigla = {"Autônomo (pessoa física)": "PF", "MEI": "MEI", "Simples Nacional": "SN", "Lucro Presumido": "LP"}[o.nome]
                lista.append(Caso(f"{cod}-{sigla}", f"Autônomo × CNPJ — {o.nome} (mensal)", desc, o.resultado))

    # Recuperação de tributos: 12 meses de histórico com a mesma receita e itens monofásicos/ST só no mês 13
    def meses_recuperacao(receita: int, **ultimo) -> list[recuperacao.MesBase]:
        return [recuperacao.MesBase(f"{2025 + i // 12}-{i % 12 + 1:02d}", Decimal(receita),
                                    **({k: Decimal(v) for k, v in ultimo.items()} if i == 12 else {}))
                for i in range(13)]

    rec = {
        "REC1": ("Comércio (Anexo I), RBT12 R$ 1.200.000; vendas de R$ 10.000 em monofásicos e R$ 20.000 com ST",
                 "I", meses_recuperacao(100_000, mono_ncm=10_000, st_nota=20_000)),
        "REC2": ("Comércio (Anexo I), RBT12 R$ 4.800.000 (6ª faixa: ICMS fora do DAS); R$ 10.000 em cada",
                 "I", meses_recuperacao(400_000, mono_nota=10_000, st_nota=10_000)),
        "REC3": ("Indústria (Anexo II), RBT12 R$ 600.000; R$ 8.000 em monofásicos e R$ 12.000 com ST",
                 "II", meses_recuperacao(50_000, mono_nota=8_000, st_nota=12_000)),
    }
    for cod, (desc, anexo, meses) in rec.items():
        r = recuperacao.estimar(meses, anexo)
        lista.append(Caso(cod, "Recuperação: PIS/Cofins monofásico e ICMS-ST (jan/2026)", desc, r.resultado))

    # Avisos de fator R: DAS no Anexo V x Anexo III e INSS do sócio sobre o pró-labore que falta (valores anuais)
    for cod, pro_labore in (("AV1", 60_000), ("AV2", 72_000)):
        p = PerfilEmpresa(Atividade.SERVICO_FATOR_R, 600_000, folha_clt_anual=96_000, pro_labore_anual=pro_labore)
        a = avisos.avaliar_fator_r(p)
        das_iii = p.receita_bruta_anual * simples.aliquota_efetiva("III", p.receita_bruta_anual, JAN_2026)
        das_v = p.receita_bruta_anual * simples.aliquota_efetiva("V", p.receita_bruta_anual, JAN_2026)
        r = Resultado("Aviso de fator R")
        r.tributos["DAS no Anexo V (fator R abaixo de 28%)"] = das_v
        r.tributos["DAS no Anexo III (fator R de 28%)"] = das_iii
        if a.codigo == "fator_r_abaixo":
            r.tributos["INSS do sócio sobre o pró-labore adicional"] = das_v - das_iii - a.valor
            r.tributos["Ganho líquido anual"] = a.valor
        else:
            r.tributos["Impacto anual se cair para o Anexo V"] = a.valor
        r.alerta(a.titulo + ". " + a.detalhe)
        desc = f"Serviços com fator R; receita R$ 600.000; folha CLT R$ 96.000; pró-labore R$ {pro_labore:,}".replace(",", ".")
        lista.append(Caso(cod, "Avisos: fator R (2026)", desc, r))
    return lista


def _cabecalho(ws, titulos: list[str], larguras: list[int]) -> None:
    for col, (t, w) in enumerate(zip(titulos, larguras), start=1):
        c = ws.cell(row=1, column=col, value=t)
        c.font = Font(name=FONTE, bold=True, color="FFFFFF")
        c.fill = CABECALHO
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


def gerar(saida: Path) -> Path:
    wb = Workbook()
    todos = casos()

    # --- Instruções ---------------------------------------------------------
    ins = wb.active
    ins.title = "Instruções"
    ins.column_dimensions["A"].width = 110
    textos = [
        ("Tolaris — validação do motor de cálculo (Fase 0)", Font(name=FONTE, bold=True, size=14, color=AZUL_ESCURO)),
        (f"Gerado em {date.today().strftime('%d/%m/%Y')} com o motor tolaris_motor. {len(todos)} casos.", None),
        ("", None),
        ("Como conferir:", Font(name=FONTE, bold=True)),
        ("1. Na aba 'Casos', cada linha é um tributo calculado pelo motor (coluna 'Valor do motor').", None),
        ("2. Preencha a coluna amarela 'Valor correto' com o valor que você apuraria.", None),
        ("3. A coluna 'Diferença' é calculada automaticamente; marque 'Confere?' e explique divergências em 'Observações'.", None),
        ("4. A aba 'Memória de cálculo' mostra cada passo, a fórmula usada e a base legal.", None),
        ("5. A aba 'Alertas' lista as premissas e limitações de cada caso.", None),
        ("", None),
        ("Premissas gerais:", Font(name=FONTE, bold=True)),
        ("• Valores anuais com faturamento mensal constante (RBT12 = receita anual no Simples).", None),
        ("• Encargos CLT no Presumido/Real: CPP 20% + RAT 2% + terceiros 5,8%; pró-labore: CPP 20%. FGTS não é tributo e fica fora do total.", None),
        ("• ICMS não cumulativo (débito - crédito sobre compras), 18% nas vendas e nas compras, salvo indicação.", None),
        ("• Lucro Real: estimativa gerencial sem LALUR detalhado; todo o estoque comprado é vendido no ano.", None),
        ("• Reforma: alíquotas de referência estimadas CBS 8,8% e IBS 17,7% (ainda não fixadas pelo Senado).", None),
        ("• Simples de 2027 a 2033 (casos SR): tabelas da Resolução CGSN 190/2026, com CBS e IBS na partilha do DAS e teto de ISS decrescente.", None),
        ("• Enquadramento (casos ENQ): cada linha é o DAS de um caminho alternativo (hoje, pró-labore até o fator R de 28%, parte da receita no Anexo III); não se somam.", None),
        ("", None),
        ("Exemplo de preenchimento: se o motor indicar R$ 28.800,00 de IRPJ e você apurar R$ 28.800,00, "
         "digite 28800 em 'Valor correto' e escolha 'Sim' em 'Confere?'.", None),
    ]
    for i, (t, f) in enumerate(textos, start=1):
        c = ins.cell(row=i, column=1, value=t)
        c.font = f or Font(name=FONTE)
        c.alignment = Alignment(wrap_text=True)

    # --- Casos --------------------------------------------------------------
    ws = wb.create_sheet("Casos")
    _cabecalho(
        ws,
        ["Caso", "Módulo", "Entradas", "Tributo", "Valor do motor (R$)", "Valor correto (R$)",
         "Diferença (R$)", "Confere?", "Observações do contador"],
        [11, 30, 70, 32, 18, 18, 16, 11, 40],
    )
    ws["E1"].comment = Comment("Valor calculado pelo tolaris_motor (Python). Não editar.", "Tolaris")
    ws["F1"].comment = Comment("Preencha com o valor que você apuraria para conferência.", "Tolaris")
    dv = DataValidation(type="list", formula1='"Sim,Não,Revisar"', allow_blank=True)
    ws.add_data_validation(dv)

    linha = 2
    for n, caso in enumerate(todos):
        fill = CINZA if n % 2 else None
        itens = list(caso.resultado.tributos.items())
        # No Simples na Reforma, no Enquadramento e nos avisos as linhas são alternativas ou resultados: somá-las não faz sentido
        if not caso.modulo.startswith(("Simples na Reforma", "Avisos", "Enquadramento")):
            itens.append(("TOTAL", None))
        inicio = linha
        for nome, valor in itens:
            ws.cell(row=linha, column=1, value=caso.codigo)
            ws.cell(row=linha, column=2, value=caso.modulo)
            ws.cell(row=linha, column=3, value=caso.descricao)
            ws.cell(row=linha, column=4, value=nome)
            if nome == "TOTAL":
                ws.cell(row=linha, column=5, value=f"=SUM(E{inicio}:E{linha - 1})" if linha > inicio else 0)
                for col in range(1, 10):
                    ws.cell(row=linha, column=col).font = Font(name=FONTE, bold=True)
            else:
                ws.cell(row=linha, column=5, value=float(round(valor, 2)))
            ws.cell(row=linha, column=6).fill = AMARELO
            ws.cell(row=linha, column=7, value=f'=IF(F{linha}="","",E{linha}-F{linha})')
            dv.add(ws.cell(row=linha, column=8))
            for col in range(1, 10):
                c = ws.cell(row=linha, column=col)
                if c.font.name != FONTE:
                    c.font = Font(name=FONTE)
                if fill and col != 6:
                    c.fill = fill
                c.border = BORDA
                c.alignment = Alignment(vertical="top", wrap_text=col in (3, 9))
                if col in (5, 6, 7):
                    c.number_format = MOEDA
            linha += 1
    ws.auto_filter.ref = f"A1:I{linha - 1}"

    # --- Memória de cálculo -------------------------------------------------
    mem = wb.create_sheet("Memória de cálculo")
    _cabecalho(mem, ["Caso", "Passo", "Fórmula", "Valor", "Base legal"], [11, 40, 90, 18, 45])
    linha = 2
    for caso in todos:
        for p in caso.resultado.memoria:
            valores = [caso.codigo, p.descricao, p.formula, float(p.valor), p.base_legal]
            for col, v in enumerate(valores, start=1):
                c = mem.cell(row=linha, column=col, value=v)
                c.font = Font(name=FONTE)
                c.alignment = Alignment(vertical="top", wrap_text=col in (2, 3, 5))
                c.border = BORDA
            eh_taxa = p.descricao.startswith(("Alíquota", "Fator R", "Percentuais"))
            mem.cell(row=linha, column=4).number_format = "0.0000%" if eh_taxa else MOEDA
            linha += 1
    mem.auto_filter.ref = f"A1:E{linha - 1}"

    # --- Alertas ------------------------------------------------------------
    al = wb.create_sheet("Alertas")
    _cabecalho(al, ["Caso", "Alerta / premissa"], [11, 140])
    linha = 2
    for caso in todos:
        for a in caso.resultado.alertas:
            for col, v in enumerate((caso.codigo, a), start=1):
                c = al.cell(row=linha, column=col, value=v)
                c.font = Font(name=FONTE)
                c.alignment = Alignment(vertical="top", wrap_text=True)
                c.border = BORDA
            linha += 1

    saida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(saida)
    return saida


if __name__ == "__main__":
    destino = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parents[1] / "validacao" / "casos_validacao.xlsx"
    print(gerar(destino))
