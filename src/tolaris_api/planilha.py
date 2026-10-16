"""Leitura de planilhas (CSV e XLSX) com receita, compras, folha e despesas por mês.

Cada sistema contábil exporta de um jeito, então o usuário escolhe quais colunas são quais; o sistema só
sugere pelo nome do cabeçalho. Tudo vira centavos por competência (AAAA-MM).
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

LIMITE_BYTES = 5 * 1024 * 1024
LIMITE_DESCOMPACTADO = 80 * 1024 * 1024
LIMITE_LINHAS = 5000
LIMITE_COLUNAS = 40
TETO_CENTAVOS = 10**13
"""R$ 100 bilhões por campo e mês: acima disso a coluna foi mapeada errado."""

CAMPOS = ("receita", "compras", "folha", "pro_labore", "despesas_creditaveis", "outras_despesas")
ROTULOS = {
    "competencia": "Competência (mês)",
    "receita": "Receita (faturamento)",
    "compras": "Compras / custo das mercadorias",
    "folha": "Folha de salários",
    "pro_labore": "Pró-labore",
    "despesas_creditaveis": "Despesas com crédito de PIS/Cofins",
    "outras_despesas": "Outras despesas",
}

SINONIMOS = {
    "competencia": ["competencia", "mes ano", "mes", "periodo", "referencia", "data"],
    "receita": ["receita bruta", "receita total", "receita", "faturamento bruto", "faturamento", "vendas"],
    "compras": ["compras", "custo das mercadorias", "custo mercadorias", "cmv", "custo", "mercadorias"],
    "folha": ["folha de pagamento", "folha clt", "folha", "salarios"],
    "pro_labore": ["pro labore", "prolabore", "pro-labore"],
    "despesas_creditaveis": ["despesas creditaveis", "despesas com credito", "despesas credito", "creditaveis"],
    "outras_despesas": ["outras despesas", "demais despesas", "despesas operacionais", "outros custos", "despesas"],
}

MESES = {"jan": 1, "fev": 2, "mar": 3, "abr": 4, "mai": 5, "jun": 6, "jul": 7, "ago": 8, "set": 9, "out": 10, "nov": 11, "dez": 12}


class PlanilhaInvalida(Exception):
    pass


def _normalizar(texto: str) -> str:
    sem_acento = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", sem_acento).strip()


# --- Leitura -----------------------------------------------------------------------------------
def _celula(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float):
        return repr(int(v)) if v == int(v) and abs(v) < 1e15 else repr(v)
    return str(v).strip()


def _ler_csv(conteudo: bytes) -> list[list[str]]:
    try:
        texto = conteudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = conteudo.decode("cp1252", errors="replace")
    primeira = next((l for l in texto.splitlines() if l.strip()), "")
    delimitador = max(";,\t", key=primeira.count)
    return [[c.strip() for c in linha] for linha in csv.reader(io.StringIO(texto), delimiter=delimitador)]


def _ler_xlsx(conteudo: bytes) -> list[list[str]]:
    from openpyxl import load_workbook  # import tardio: só quem sobe XLSX paga o custo

    try:
        with zipfile.ZipFile(io.BytesIO(conteudo)) as z:
            if sum(i.file_size for i in z.infolist()) > LIMITE_DESCOMPACTADO:
                raise PlanilhaInvalida("Planilha grande demais depois de descompactada")
        wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except zipfile.BadZipFile:
        raise PlanilhaInvalida("O arquivo não é uma planilha .xlsx válida") from None
    except PlanilhaInvalida:
        raise
    except Exception:  # noqa: BLE001 - qualquer falha do leitor vira mensagem para o usuário
        raise PlanilhaInvalida("Não foi possível ler a planilha .xlsx") from None
    try:
        linhas = []
        for linha in wb.active.iter_rows(values_only=True, max_col=LIMITE_COLUNAS):
            celulas = [_celula(c) for c in linha]
            while celulas and not celulas[-1]:  # o leitor completa até a largura máxima com vazios
                celulas.pop()
            linhas.append(celulas)
            if len(linhas) > LIMITE_LINHAS + 1:
                break
        return linhas
    finally:
        wb.close()


def ler_tabela(nome: str, conteudo: bytes) -> tuple[list[str], list[list[str]]]:
    """(cabeçalhos, linhas de dados) de um .csv ou .xlsx; a primeira linha não vazia é o cabeçalho."""
    if len(conteudo) > LIMITE_BYTES:
        raise PlanilhaInvalida(f"Arquivo acima de {LIMITE_BYTES // 1024 // 1024} MB")
    minusculo = nome.lower()
    if minusculo.endswith((".csv", ".txt")):
        bruto = _ler_csv(conteudo)
    elif minusculo.endswith(".xlsx"):
        bruto = _ler_xlsx(conteudo)
    else:
        raise PlanilhaInvalida("Envie um arquivo .csv ou .xlsx")
    linhas = [[c for c in l[:LIMITE_COLUNAS]] for l in bruto if any(c.strip() for c in l)]
    if len(linhas) < 2:
        raise PlanilhaInvalida("A planilha precisa ter o cabeçalho e ao menos uma linha de dados")
    if len(linhas) - 1 > LIMITE_LINHAS:
        raise PlanilhaInvalida(f"Mais de {LIMITE_LINHAS} linhas; divida a planilha")
    largura = max(len(l) for l in linhas)
    linhas = [l + [""] * (largura - len(l)) for l in linhas]
    return linhas[0], linhas[1:]


# --- Mapeamento de colunas -------------------------------------------------------------------------
def sugerir_mapeamento(
    cabecalhos: list[str],
    sinonimos: dict[str, list[str]] = SINONIMOS,
    campos: tuple[str, ...] = ("competencia", *CAMPOS),
) -> dict[str, int | None]:
    """Índice da coluna sugerida para cada campo (ou None), pelo nome do cabeçalho."""
    normalizados = [_normalizar(c) for c in cabecalhos]
    usadas: set[int] = set()
    sugestao: dict[str, int | None] = {}
    for campo in campos:
        achada = None
        for sinonimo in sinonimos[campo]:  # do mais específico ao mais geral
            alvo = _normalizar(sinonimo)
            exata = next((i for i, c in enumerate(normalizados) if c == alvo and i not in usadas), None)
            parcial = next((i for i, c in enumerate(normalizados) if alvo in c.split() or (len(alvo) > 5 and alvo in c)
                            if i not in usadas), None) if exata is None else None
            achada = exata if exata is not None else parcial
            if achada is not None:
                break
        if achada is not None:
            usadas.add(achada)
        sugestao[campo] = achada
    return sugestao


# --- Conversão de valores -----------------------------------------------------------------------------
def parse_competencia(texto: str) -> str | None:
    t = texto.strip().lower()
    if not t:
        return None
    t = unicodedata.normalize("NFKD", t).encode("ascii", "ignore").decode()
    m = re.fullmatch(r"(\d{4})[-/.](\d{1,2})(?:[-/.]\d{1,2})?(?:[ t].*)?", t)  # 2026-03, 2026-03-01, 2026/03
    if m:
        ano, mes = int(m[1]), int(m[2])
    else:
        m = re.fullmatch(r"(?:\d{1,2}[/.-])?(\d{1,2})[/.-](\d{4})", t)  # 03/2026, 01/03/2026
        if m:
            mes, ano = int(m[1]), int(m[2])
        else:
            m = re.fullmatch(r"([a-z]{3})[a-z]*[/ .-]+(\d{2}|\d{4})", t)  # mar/26, março/2026
            if not m or m[1] not in MESES:
                return None
            mes, ano = MESES[m[1]], int(m[2])
            ano += 2000 if ano < 100 else 0
    return f"{ano:04d}-{mes:02d}" if 1 <= mes <= 12 and 2000 <= ano <= 2100 else None


def parse_dinheiro(texto: str) -> int | None:
    """Reais em texto (1.234,56 | 1234.56 | R$ 1.234,56) para centavos; vazio vira None; inválido levanta ValueError."""
    t = re.sub(r"(?i)r\$|\s", "", texto)
    if not t or t in ("-", "–"):
        return None
    if t.startswith("(") and t.endswith(")"):
        raise ValueError("valor negativo")
    t = t.removeprefix("+")
    if t.startswith("-"):
        raise ValueError("valor negativo")
    if "," in t and "." in t:
        t = t.replace(".", "").replace(",", ".") if t.rfind(",") > t.rfind(".") else t.replace(",", "")
    elif "," in t:
        t = t.replace(",", ".")
    elif "." in t and re.fullmatch(r"\d{1,3}(\.\d{3})+", t):
        t = t.replace(".", "")  # 1.234 = mil duzentos e trinta e quatro
    try:
        valor = Decimal(t)
    except InvalidOperation:
        raise ValueError(f"valor inválido: {texto[:20]}") from None
    centavos = int((valor * 100).quantize(Decimal(1)))
    if centavos >= TETO_CENTAVOS:
        raise ValueError("valor grande demais; confira a coluna escolhida")
    return centavos


@dataclass
class ResultadoConversao:
    meses: dict[str, dict[str, int | None]] = field(default_factory=dict)
    ignoradas: list[tuple[int, str]] = field(default_factory=list)
    """(número da linha na planilha, motivo)."""
    avisos: list[str] = field(default_factory=list)


def validar_mapeamento(mapeamento: dict[str, int | None], colunas: int) -> dict[str, int]:
    limpo = {k: v for k, v in mapeamento.items() if v is not None and k in ("competencia", *CAMPOS)}
    if "competencia" not in limpo:
        raise PlanilhaInvalida("Indique a coluna da competência (o mês)")
    if not any(c in limpo for c in CAMPOS):
        raise PlanilhaInvalida("Indique ao menos uma coluna de valores (receita, folha, despesas...)")
    if any(not isinstance(i, int) or i < 0 or i >= colunas for i in limpo.values()):
        raise PlanilhaInvalida("Coluna escolhida fora da planilha")
    if len(set(limpo.values())) != len(limpo):
        raise PlanilhaInvalida("A mesma coluna foi escolhida para dois campos")
    return limpo


def converter(linhas: list[list[str]], mapeamento: dict[str, int]) -> ResultadoConversao:
    """Converte as linhas em valores por competência; linhas com problema são ignoradas e explicadas."""
    res = ResultadoConversao()
    repetidas: set[str] = set()
    for n, linha in enumerate(linhas, start=2):  # linha 1 é o cabeçalho
        comp = parse_competencia(linha[mapeamento["competencia"]])
        if comp is None:
            res.ignoradas.append((n, f"competência inválida: “{linha[mapeamento['competencia']][:20]}”"))
            continue
        try:
            valores = {c: parse_dinheiro(linha[i]) for c, i in mapeamento.items() if c != "competencia"}
        except ValueError as erro:
            res.ignoradas.append((n, str(erro)))
            continue
        atual = res.meses.setdefault(comp, {c: None for c in mapeamento if c != "competencia"})
        if any(v is not None for v in atual.values()):
            repetidas.add(comp)
        for c, v in valores.items():
            if v is not None:
                atual[c] = (atual[c] or 0) + v
    if repetidas:
        res.avisos.append("Competências com mais de uma linha tiveram os valores somados: " + ", ".join(sorted(repetidas)) + ".")
    return res


def modelo_csv() -> bytes:
    """Planilha de exemplo para o contador preencher (UTF-8 com BOM e ponto e vírgula, que o Excel brasileiro abre direto)."""
    linhas = [
        ["Competência", "Receita", "Compras", "Folha", "Pró-labore", "Despesas creditáveis", "Outras despesas"],
        ["2026-01", "120000,00", "72000,00", "18000,00", "6000,00", "3500,00", "9000,00"],
        ["2026-02", "115000,00", "69000,00", "18000,00", "6000,00", "3400,00", "8800,00"],
    ]
    saida = io.StringIO()
    csv.writer(saida, delimiter=";", lineterminator="\r\n").writerows(linhas)
    return ("﻿" + saida.getvalue()).encode("utf-8")
